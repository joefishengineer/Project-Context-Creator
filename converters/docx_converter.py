"""
DOCX to Markdown Converter
Extracts text from Word documents
"""

from pathlib import Path
import logging

def convert_docx_to_markdown(docx_path, output_path=None):
    """
    Convert DOCX to markdown, extracting all text content.
    
    Args:
        docx_path: Path to DOCX file
        output_path: Optional output path (defaults to same name with .md extension)
    
    Returns:
        Path to created markdown file
    """
    try:
        from docx import Document
    except ImportError:
        raise ImportError(
            "Required library not installed. Run setup.ps1 or install manually:\n"
            "pip install python-docx"
        )
    
    docx_path = Path(docx_path)
    if output_path is None:
        output_path = docx_path.with_suffix('.md')
    else:
        output_path = Path(output_path)
    
    logging.info(f"Converting DOCX: {docx_path}")
    
    try:
        doc = Document(str(docx_path))
    except Exception as e:
        logging.error(f"Failed to open DOCX file: {e}")
        raise
    
    # Extract text
    markdown_lines = []
    markdown_lines.append(f"# {docx_path.stem}\n")
    markdown_lines.append(f"*Converted from: {docx_path.name}*\n\n")
    
    for para in doc.paragraphs:
        try:
            text = para.text.strip()
            if text:
                # Simple heading detection based on style
                try:
                    if para.style and para.style.name and para.style.name.startswith('Heading'):
                        level = para.style.name.replace('Heading ', '').strip()
                        if level.isdigit():
                            markdown_lines.append(f"\n{'#' * int(level)} {text}\n")
                        else:
                            markdown_lines.append(f"\n## {text}\n")
                    else:
                        markdown_lines.append(f"{text}\n")
                except (AttributeError, TypeError) as e:
                    # Style is None or malformed - treat as regular text
                    logging.warning(f"Paragraph has invalid style: {e}")
                    markdown_lines.append(f"{text}\n")
        except Exception as e:
            logging.warning(f"Could not extract paragraph text: {e}")
            continue
    
    # Extract tables
    if doc.tables:
        markdown_lines.append("\n---\n\n## Tables\n")
        for i, table in enumerate(doc.tables, 1):
            try:
                markdown_lines.append(f"\n### Table {i}\n\n")
                for row in table.rows:
                    try:
                        # Deduplicate merged cells: python-docx intentionally returns
                        # the same cell object multiple times for each grid column a
                        # horizontally-merged cell spans (documented in _Row.cells).
                        # Comparing ._tc identity gives us only the unique XML elements.
                        seen_tcs = set()
                        unique_cells = []
                        for cell in row.cells:
                            if id(cell._tc) not in seen_tcs:
                                seen_tcs.add(id(cell._tc))
                                unique_cells.append(cell)
                        row_text = " | ".join(cell.text.strip() for cell in unique_cells)
                        if row_text.strip().replace("|", "").strip():
                            markdown_lines.append(f" | {row_text} |\n")
                    except Exception as e:
                        logging.warning(f"Table {i}: Could not extract row: {e}")
                        continue
                markdown_lines.append("\n")
            except Exception as e:
                logging.warning(f"Could not extract table {i}: {e}")
                continue
    
    # Write markdown file
    content = '\n'.join(markdown_lines)
    output_path.write_text(content, encoding='utf-8')
    
    logging.info(f"Created markdown: {output_path}")
    return output_path

# Made with Bob
